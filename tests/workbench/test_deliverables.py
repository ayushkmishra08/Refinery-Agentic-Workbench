"""Word / PowerPoint / Excel deliverables built from real (LLM-free, mock-backend) responses."""
from __future__ import annotations

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from workbench.core.request import UserRequest
from workbench.deliverables.common import calculation_steps_from_response, figures_from_response
from workbench.deliverables.export import export_markdown, export_response


@pytest.fixture(scope="module")
def lookup_resp(orch):
    return orch.ask(UserRequest(text="What is the normal flow rate of the crude charge pump?", session_id="deliv-lookup"))


@pytest.fixture(scope="module")
def limits_resp(orch):
    return orch.ask(UserRequest(text="The crude charge pump is operating at 520 m3/h. Is this acceptable?", session_id="deliv-limits"))


def test_figures_and_steps_are_extracted(lookup_resp, limits_resp):
    figs = figures_from_response(lookup_resp)
    assert figs, "a lookup answer states at least one figure"
    assert any(f["citation"] for f in figs)
    assert any(f["document_id"] for f in figs), "figures join to their evidence"
    steps = calculation_steps_from_response(limits_resp)
    assert steps and steps[-1]["kind"] == "verdict"
    assert any(s["kind"] == "arith" for s in steps), "the deviation is shown as arithmetic"


def test_docx_export(lookup_resp, tmp_path):
    out = export_response(lookup_resp, "docx", tmp_path)
    assert out.fmt == "docx" and out.bytes > 0 and len(out.sha256) == 64
    doc = Document(out.path)
    header = doc.sections[0].header.paragraphs[0].text
    assert "pending human sign-off" in header
    assert (lookup_resp.security.classification or "UNCLASSIFIED") in header
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "pending human sign-off" in body
    assert any(p.text == "Provenance" for p in doc.paragraphs if p.style.name.startswith("Heading"))
    assert len(doc.tables) >= 1
    assert doc.core_properties.author == "Refinery Engineering AI Workbench"


def test_xlsx_export_has_live_formula_for_limits(limits_resp, lookup_resp, tmp_path):
    out = export_response(limits_resp, "xlsx", tmp_path)
    wb = load_workbook(out.path)
    for name in ("Answer", "Values", "Calculation", "Provenance", "Evidence", "Audit"):
        assert name in wb.sheetnames, name
    calc = wb["Calculation"]
    formulas = [c.value for row in calc.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert formulas, "the calculation sheet recomputes with a real formula"
    assert calc["A1"].value == "Step" and calc.freeze_panes == "A2"
    values = wb["Values"]
    assert values.max_row >= 2
    out2 = export_response(lookup_resp, "xlsx", tmp_path)
    assert load_workbook(out2.path)["Answer"]["A1"].value


def test_pptx_export(limits_resp, tmp_path):
    out = export_response(limits_resp, "pptx", tmp_path)
    prs = Presentation(out.path)
    slides = list(prs.slides)
    assert len(slides) >= 4
    titles = [s.shapes.title.text for s in slides if s.shapes.title is not None]
    assert "Answer" in titles and "Provenance" in titles
    last_texts = [sh.text_frame.text for sh in slides[-1].shapes if sh.has_text_frame]
    assert any("pending human sign-off" in t for t in last_texts)


def test_markdown_export_and_bad_format(lookup_resp, tmp_path):
    out = export_markdown(lookup_resp, tmp_path)
    text = open(out.path, encoding="utf-8").read()
    assert "## Provenance" in text and "## Evidence" in text
    with pytest.raises(ValueError):
        export_response(lookup_resp, "pdfx", tmp_path)
