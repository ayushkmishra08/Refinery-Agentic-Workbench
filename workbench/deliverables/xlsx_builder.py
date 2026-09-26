"""Excel deliverable (openpyxl). Calculation steps become rows with real formulas where the
step is arithmetic over values already on the sheet, so the reviewer can recompute in place."""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from workbench.agents.base import REASONING_BLOCKS
from workbench.deliverables.common import (
    DeliverableSpec, audit_line, calculation_steps_from_response, evidence_rows, figures_from_response,
    markdown_to_paragraphs, strip_banner,
)

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
RED = Font(bold=True, color="B01E1E")
_SHEET_BAD = re.compile(r"[\[\]:*?/\\]")


def _num(v):
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip().replace(",", ".")
    try:
        return float(s) if re.fullmatch(r"-?\d+(?:\.\d+)?", s) else v
    except ValueError:
        return v


def _sheet(wb: Workbook, name: str, columns: list[str], rows: list[list], *, widths: dict[int, int] | None = None):
    title = _SHEET_BAD.sub("-", name)[:31] or "Sheet"
    base, n = title, 2
    while title in wb.sheetnames:
        title = f"{base[:28]}_{n}"
        n += 1
    ws = wb.create_sheet(title)
    ws.append(columns)
    for c in ws[1]:
        c.font, c.fill = HEAD, HEAD_FILL
    for row in rows:
        ws.append([_num(v) if v is not None else None for v in row])
    ws.freeze_panes = "A2"
    for i, col in enumerate(columns, start=1):
        width = (widths or {}).get(i)
        if width is None:
            longest = max([len(str(col))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r) and r[i - 1] is not None] or [8])
            width = min(60, max(10, longest + 2))
        ws.column_dimensions[get_column_letter(i)].width = width
    return ws


def build_xlsx(resp, spec: DeliverableSpec, out_path: Path, figures=None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Answer"
    ws["A1"] = strip_banner(resp.answer_markdown)
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 110
    ws.row_dimensions[1].height = 220
    meta = [("Title", spec.title), ("Classification", spec.classification or "UNCLASSIFIED"), ("Status", spec.status_line),
            ("Author", spec.author), ("Generated", spec.generated_at), ("Response", getattr(resp, "response_id", "")),
            ("Audit", getattr(resp, "audit_trail_id", "")), ("Confidence", f"{resp.confidence.score:.2f} ({resp.confidence.level})"),
            ("Human review", "required" if getattr(resp, "requires_human_review", False) else "not flagged")]
    for i, (k, v) in enumerate(meta, start=3):
        ws.cell(row=i, column=1, value=f"{k}: {v}")
        if k == "Status" and not spec.signed_off:
            ws.cell(row=i, column=1).font = RED
    wb.properties.creator = spec.author
    wb.properties.title = spec.title
    wb.properties.category = spec.classification or ""

    figs = figures if figures is not None else figures_from_response(resp)
    fig_rows = []
    for f in figs:
        d = f if isinstance(f, dict) else f.model_dump()
        src = d.get("source") or {}
        fig_rows.append([d.get("label", ""), d.get("value", ""), d.get("unit") or "", d.get("qualifier") or "", d.get("citation") or "",
                         d.get("document_id") or src.get("document_id") or "", d.get("page") or src.get("page") or ""])
    values_ws = _sheet(wb, "Values", ["Label", "Value", "Unit", "Qualifier", "Ref", "Document", "Page"], fig_rows)

    for b in resp.blocks:
        t = b.type
        if t in REASONING_BLOCKS or t == "evidence":
            continue
        if t == "table":
            cites = b.row_citations or []
            _sheet(wb, b.title or "Table", [*b.columns, "Ref"], [[*r, ", ".join(cites[i]) if i < len(cites) else ""] for i, r in enumerate(b.rows)])
        elif t == "comparison":
            _sheet(wb, b.title or "Comparison", ["Attribute", *[f"{s}" for s in b.subjects], "Refs"],
                   [[a, *[f"{c.value or ''} {c.unit or ''}".strip() for c in row], ", ".join(c.citation for c in row if c.citation)] for a, row in zip(b.attributes, b.cells)])
        elif t == "kpi":
            _sheet(wb, b.title or "Documented values", ["Label", "Value", "Unit", "Qualifier", "Ref"],
                   [[i.label, i.value, i.unit or "", i.qualifier or "", i.citation or ""] for i in b.items])
        elif t == "conflict":
            _sheet(wb, "Conflict", ["Value", "Unit", "Document", "Rev", "Page", "Source", "Context", "Ref", "Preferred"],
                   [[c.value, c.unit or "", c.document_id, c.revision or "", c.page, c.source, c.context or "", c.citation or "", "yes" if b.preferred_index == i else ""] for i, c in enumerate(b.claims)])
        elif t == "steps":
            rows = [[0, "prerequisite", p.text, p.page or "", p.citation or "", ""] for p in b.prerequisites]
            rows += [[s.sequence, "step", s.text, s.page or "", s.citation or "", "; ".join(s.warnings)] for s in b.steps]
            _sheet(wb, "Steps", ["#", "Kind", "Text", "Page", "Ref", "Warnings"], rows, widths={3: 90})
        elif t == "safety":
            _sheet(wb, "Safety", ["Severity", "Message", "Ref", "Needs authorisation"],
                   [[f.severity, f.message, f.citation or "", "yes" if f.requires_authorization else ""] for f in b.flags], widths={2: 90})
        elif t == "graph":
            names = {n.id: n.label for n in b.nodes}
            _sheet(wb, "Topology", ["From", "Relation", "To", "Ref", "Inferred"],
                   [[names.get(e.source, e.source), e.label, names.get(e.target, e.target), e.citation or "", "yes" if e.inferred else ""] for e in b.edges])

    # ---- calculation sheet with live formulas ------------------------------------------
    steps = calculation_steps_from_response(resp)
    calc = wb.create_sheet("Calculation")
    calc.append(["Step", "Expression", "Literal result", "Recomputed (formula)", "Note", "Ref"])
    for c in calc[1]:
        c.font, c.fill = HEAD, HEAD_FILL
    calc.freeze_panes = "A2"
    value_cells: dict[str, str] = {}      # marker label -> cell ref holding the number
    for s in steps:
        r = calc.max_row + 1
        calc.cell(row=r, column=1, value=s["step"])
        calc.cell(row=r, column=2, value=s["expression"])
        calc.cell(row=r, column=5, value=s["note"])
        calc.cell(row=r, column=6, value=", ".join(x for x in s["refs"] if x))
        if s["kind"] == "input":
            m = re.match(r"(-?\d+(?:\.\d+)?)", s["result"])
            calc.cell(row=r, column=3, value=float(m.group(1)) if m else s["result"])
            label = s["expression"].split(" ")[0].lower()
            value_cells[label] = f"C{r}"
        elif s["kind"] == "arith" and s.get("op") == "pct_dev":
            calc.cell(row=r, column=3, value=s["result"])
            v, n = value_cells.get("requested"), (value_cells.get("normal") or value_cells.get("rated") or value_cells.get("operating"))
            if v and n:
                calc.cell(row=r, column=4, value=f"=({v}-{n})/{n}*100")
        elif s["kind"] == "check":
            calc.cell(row=r, column=3, value=s["result"])
            lab = re.search(r"\((\w+)\)", s["expression"])
            v, lim = value_cells.get("requested"), value_cells.get(lab.group(1).lower()) if lab else None
            if v and lim:
                calc.cell(row=r, column=4, value=f"={v}<={lim}" if s.get("op") == "le" else f"={v}>={lim}")
        else:
            calc.cell(row=r, column=3, value=s["result"])
    if not steps:
        calc.append(["—", "no calculation in this answer", "", "", "", ""])
    for i, w in enumerate((6, 44, 22, 24, 40, 10), start=1):
        calc.column_dimensions[get_column_letter(i)].width = w

    prov = []
    for f in figs:
        d = f if isinstance(f, dict) else f.model_dump()
        src = d.get("source") or {}
        prov.append([d.get("label", ""), d.get("value", ""), d.get("unit") or "", d.get("document_id") or src.get("document_id") or "",
                     d.get("page") or src.get("page") or "", d.get("chunk_id") or src.get("chunk_id") or "", d.get("claim_id") or src.get("claim_id") or "",
                     (d.get("source_text") or src.get("text") or "")[:200], d.get("confidence", 1.0) if not isinstance(f, dict) else "",
                     ("FLAGGED: " + (d.get("flag_reason") or "")) if d.get("flagged") else ""])
    _sheet(wb, "Provenance", ["Figure", "Value", "Unit", "Document", "Page", "Chunk", "Claim", "Source excerpt", "Confidence", "Flag"], prov, widths={8: 70})
    _sheet(wb, "Evidence", ["Ref", "Document", "Page", "Source", "Section", "Text"],
           [[e["ref"], e["document_id"], e["page"], e["source"], e["section"] or "", e["text"][:500]] for e in evidence_rows(resp)], widths={6: 100})
    _sheet(wb, "Audit", ["Field", "Value"], [["Summary", audit_line(resp)], ["Status", spec.status_line], ["Classification", spec.classification or ""],
                                             ["Task", getattr(getattr(resp, "task_type", None), "value", "")], ["Timing ms", getattr(resp, "timing_ms", 0)]], widths={2: 120})
    values_ws.sheet_properties.tabColor = "1F3A5F"
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    return out_path
