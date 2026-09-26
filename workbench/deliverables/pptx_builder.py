"""PowerPoint deliverable (python-pptx)."""
from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from workbench.agents.base import REASONING_BLOCKS
from workbench.deliverables.common import (
    DeliverableSpec, audit_line, calculation_steps_from_response, evidence_rows, figures_from_response,
    markdown_to_paragraphs, strip_banner,
)

RED = RGBColor(0xB0, 0x1E, 0x1E)
GREY = RGBColor(0x55, 0x55, 0x55)
W, H = Inches(13.333), Inches(7.5)


def _footer(slide, spec: DeliverableSpec) -> None:
    tb = slide.shapes.add_textbox(Inches(0.4), H - Inches(0.45), W - Inches(0.8), Inches(0.35))
    p = tb.text_frame.paragraphs[0]
    p.text = f"{spec.classification or 'UNCLASSIFIED'} · {spec.status_line} · {spec.author}"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = RED if not spec.signed_off else GREY


def _title_only(prs, title: str, spec: DeliverableSpec):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(28)
    _footer(slide, spec)
    return slide


def _bullets(prs, title: str, lines: list[str], spec: DeliverableSpec, *, cap: int = 12, red_title: bool = False):
    chunks = [lines[i:i + cap] for i in range(0, len(lines), cap)] or [[]]
    for n, chunk in enumerate(chunks):
        slide = _title_only(prs, title + (f" ({n + 1}/{len(chunks)})" if len(chunks) > 1 else ""), spec)
        if red_title:
            slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = RED
        tb = slide.shapes.add_textbox(Inches(0.6), Inches(1.5), W - Inches(1.2), H - Inches(2.2))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, line in enumerate(chunk):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"• {line}"
            p.font.size = Pt(16 if len(chunk) <= 6 else 13)


def _table_slide(prs, title: str, columns: list[str], rows: list[list], spec: DeliverableSpec, *, cap: int = 12):
    rows = rows[:cap]
    slide = _title_only(prs, title, spec)
    n_rows, n_cols = len(rows) + 1, max(1, len(columns))
    shape = slide.shapes.add_table(n_rows, n_cols, Inches(0.5), Inches(1.4), W - Inches(1.0), Inches(0.4) * n_rows)
    table = shape.table
    for j, c in enumerate(columns):
        cell = table.cell(0, j)
        cell.text = str(c)
        cell.text_frame.paragraphs[0].font.size = Pt(12)
        cell.text_frame.paragraphs[0].font.bold = True
    for i, row in enumerate(rows, start=1):
        for j in range(n_cols):
            v = row[j] if j < len(row) else ""
            cell = table.cell(i, j)
            cell.text = "" if v is None else str(v)[:90]
            cell.text_frame.paragraphs[0].font.size = Pt(11)
    return slide


def _answer_bullets(md: str, max_bullets: int = 6, max_words: int = 25) -> list[str]:
    text = " ".join(t for s, t in markdown_to_paragraphs(strip_banner(md)) if s != "table_row")
    sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", text) if s.strip()]
    out: list[str] = []
    for s in sentences:
        words = s.split()
        out.append(" ".join(words[:max_words]) + (" …" if len(words) > max_words else ""))
        if len(out) >= max_bullets:
            break
    return out or ["(no prose answer)"]


def build_pptx(resp, spec: DeliverableSpec, out_path: Path, figures=None) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    # title slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = spec.title
    sub = slide.placeholders[1]
    sub.text = f"{spec.subtitle}\n{spec.classification or 'UNCLASSIFIED'} · {spec.status_line}\n{spec.generated_at}"
    for p in sub.text_frame.paragraphs:
        p.font.size = Pt(16)
    if not spec.signed_off:
        sub.text_frame.paragraphs[1].font.color.rgb = RED
        sub.text_frame.paragraphs[1].font.bold = True
    _footer(slide, spec)

    _bullets(prs, "Answer", _answer_bullets(resp.answer_markdown), spec, cap=6)

    for b in resp.blocks:
        t = b.type
        if t in REASONING_BLOCKS or t in ("evidence", "text"):
            continue
        if t == "callout":
            _bullets(prs, b.title or "Note", [x for _, x in markdown_to_paragraphs(b.markdown)], spec, red_title=b.level in ("danger", "warning"))
        elif t == "kpi":
            _table_slide(prs, b.title or "Documented values", ["Value", "Figure"],
                         [[i.label, f"{i.value} {i.unit or ''} {('(' + i.qualifier + ')') if i.qualifier else ''} {i.citation or ''}".strip()] for i in b.items], spec)
        elif t == "table":
            cites = b.row_citations or []
            _table_slide(prs, b.title or "Table", [*b.columns, "Ref"],
                         [[*r, ", ".join(cites[i]) if i < len(cites) else ""] for i, r in enumerate(b.rows)], spec)
        elif t == "steps":
            lines = [f"[prerequisite] {p.text} {p.citation or ''}".strip() for p in b.prerequisites]
            lines += [f"{s.sequence}. {s.text} {s.citation or ''}".strip() + (f"  ⚠ {'; '.join(s.warnings)}" if s.warnings else "") for s in b.steps]
            _bullets(prs, b.title or "Procedure", lines, spec, cap=12)
        elif t == "safety":
            _bullets(prs, b.title or "Safety", [f"{f.severity.upper()}: {f.message} {f.citation or ''}".strip() for f in b.flags], spec, red_title=True)
        elif t == "limit_gauge":
            s = _table_slide(prs, b.title or f"Limit check — {b.entity}", ["Marker", "Value", "Unit", "Ref"],
                             [[m.label, f"{m.value:g}", b.unit, m.citation or ""] for m in b.markers] + [["checked value", f"{b.value:g}" if b.value is not None else "—", b.unit, ""]], spec)
            tb = s.shapes.add_textbox(Inches(0.5), H - Inches(1.4), W - Inches(1.0), Inches(0.8))
            tb.text_frame.text = f"Verdict: {b.verdict.replace('_', ' ')}. {b.message}"
            tb.text_frame.paragraphs[0].font.size = Pt(14)
            tb.text_frame.paragraphs[0].font.bold = True
        elif t == "comparison":
            _table_slide(prs, b.title or "Comparison", ["Attribute", *b.subjects],
                         [[a, *[f"{c.value or '—'} {c.unit or ''} {c.citation or ''}".strip() for c in row]] for a, row in zip(b.attributes, b.cells)], spec)
        elif t == "conflict":
            _table_slide(prs, b.title or f"Documented values — {b.subject}", ["Value", "Unit", "Document", "Page", "Source", "Context", "Ref"],
                         [[c.value, c.unit or "", c.document_id, c.page or "", c.source, c.context or "", c.citation or ""] for c in b.claims], spec)
        elif t == "graph":
            names = {n.id: n.label for n in b.nodes}
            _table_slide(prs, b.title or "Topology", ["From", "Relation", "To", "Ref"],
                         [[names.get(e.source, e.source), e.label, names.get(e.target, e.target), e.citation or ""] for e in b.edges], spec)
        elif t == "clarification":
            _bullets(prs, "Clarification needed", [b.question, *[f"missing: {m}" for m in b.missing]], spec)

    steps = calculation_steps_from_response(resp)
    if steps:
        _table_slide(prs, "Calculation steps", ["Step", "Expression", "Result", "Note", "Ref"],
                     [[s["step"], s["expression"], s["result"], s["note"], ", ".join(r for r in s["refs"] if r)] for s in steps], spec, cap=14)

    figs = figures if figures is not None else figures_from_response(resp)
    rows = []
    for f in figs:
        d = f if isinstance(f, dict) else f.model_dump()
        src = d.get("source") or {}
        rows.append([d.get("label", ""), f"{d.get('value', '')} {d.get('unit') or ''}".strip(),
                     d.get("document_id") or src.get("document_id") or "—", d.get("page") or src.get("page") or "—",
                     ("FLAGGED: " + (d.get("flag_reason") or "")) if d.get("flagged") else (d.get("citation") or "")])
    _table_slide(prs, "Provenance", ["Figure", "Value", "Document", "Page", "Ref / flag"], rows or [["—", "", "", "", "no figures"]], spec, cap=14)

    _table_slide(prs, "Evidence", ["Ref", "Document", "Page", "Source", "Text"],
                 [[e["ref"], e["document_id"], e["page"], e["source"], e["text"][:80]] for e in evidence_rows(resp)] or [["—", "", "", "", "no evidence"]], spec, cap=10)

    last = _title_only(prs, "Audit", spec)
    tb = last.shapes.add_textbox(Inches(0.6), Inches(1.6), W - Inches(1.2), Inches(2))
    tb.text_frame.word_wrap = True
    tb.text_frame.text = audit_line(resp)
    tb.text_frame.paragraphs[0].font.size = Pt(12)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path
