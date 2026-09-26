"""Word deliverable (python-docx)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from workbench.agents.base import REASONING_BLOCKS
from workbench.deliverables.common import (
    DeliverableSpec, audit_line, calculation_steps_from_response, evidence_rows, figures_from_response,
    markdown_to_paragraphs, strip_banner,
)

RED = RGBColor(0xB0, 0x1E, 0x1E)
GREY = RGBColor(0x55, 0x55, 0x55)


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _table(doc, columns: list[str], rows: list[list], style: str = "Light Grid Accent 1"):
    table = doc.add_table(rows=1, cols=len(columns))
    try:
        table.style = style
    except Exception:
        table.style = "Table Grid"
    for i, c in enumerate(columns):
        cell = table.rows[0].cells[i]
        cell.text = str(c)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row[: len(columns)]):
            cells[i].text = "" if v is None else str(v)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    return table


def _heading(doc, text: str, level: int = 1):
    return doc.add_heading(text, level=level)


def _para(doc, text: str, *, bold=False, italic=False, color=None, size=None, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    if size:
        run.font.size = Pt(size)
    return p


def _render_block(doc, b) -> None:
    t = b.type
    if t in REASONING_BLOCKS or t == "evidence":
        return
    if t == "text":
        if b.title:
            _heading(doc, b.title, 2)
        for style, text in markdown_to_paragraphs(b.markdown):
            _md_para(doc, style, text)
    elif t == "callout":
        _heading(doc, b.title or {"danger": "Warning", "warning": "Caution", "success": "Note", "info": "Note"}[b.level], 2)
        table = doc.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        cell = table.rows[0].cells[0]
        _shade(cell, {"danger": "F8D7DA", "warning": "FFF3CD", "success": "D4EDDA", "info": "E2E8F0"}[b.level])
        cell.text = ""
        for style, text in markdown_to_paragraphs(b.markdown):
            p = cell.add_paragraph()
            p.add_run(("• " if style == "bullet" else "") + text).font.size = Pt(9.5)
    elif t == "kpi":
        _heading(doc, b.title or "Documented values", 2)
        _table(doc, ["Value", "Figure", "Unit", "Qualifier", "Ref"],
               [[i.label, i.value, i.unit or "", i.qualifier or "", i.citation or ""] for i in b.items])
    elif t == "table":
        _heading(doc, b.title or "Table", 2)
        cites = b.row_citations or []
        rows = [[*r, ", ".join(cites[i]) if i < len(cites) else ""] for i, r in enumerate(b.rows)]
        _table(doc, [*b.columns, "Ref"], rows)
        if b.caption:
            _para(doc, b.caption, italic=True, size=9, color=GREY)
    elif t == "steps":
        _heading(doc, b.title or "Procedure", 2)
        meta = " · ".join(x for x in (b.document_id, b.section_path, f"p.{b.page_start}" if b.page_start else None) if x)
        if meta:
            _para(doc, meta, italic=True, size=9, color=GREY)
        if b.prerequisites:
            _para(doc, "Prerequisites", bold=True)
            for p in b.prerequisites:
                _para(doc, f"{p.text} {p.citation or ''}".strip(), style="List Bullet")
        for s in b.steps:
            para = doc.add_paragraph(style="List Number")
            para.add_run(f"{s.text} ")
            if s.citation:
                r = para.add_run(s.citation)
                r.font.superscript = True
            if s.warnings:
                w = para.add_run("  ⚠ " + "; ".join(s.warnings))
                w.font.color.rgb = RED
                w.font.size = Pt(9)
    elif t == "safety":
        _heading(doc, b.title or "Safety", 2)
        table = doc.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        cell = table.rows[0].cells[0]
        _shade(cell, "FFF3CD")
        cell.text = ""
        for f in b.flags:
            p = cell.add_paragraph()
            r = p.add_run(f"{f.severity.upper()}: ")
            r.bold = True
            r.font.color.rgb = RED if f.severity in ("danger", "warning") else GREY
            p.add_run(f"{f.message} {f.citation or ''}".strip()).font.size = Pt(9.5)
    elif t == "comparison":
        _heading(doc, b.title or "Comparison", 2)
        rows = [[a, *[f"{c.value or '—'} {c.unit or ''} {c.citation or ''}".strip() for c in row]] for a, row in zip(b.attributes, b.cells)]
        _table(doc, ["Attribute", *b.subjects], rows)
        for d in b.differences:
            _para(doc, d, style="List Bullet")
    elif t == "limit_gauge":
        _heading(doc, b.title or "Limit check", 2)
        _para(doc, f"{b.entity} — {b.parameter}: {b.value if b.value is not None else '—'} {b.unit} → {b.verdict.replace('_', ' ')}", bold=True)
        _table(doc, ["Marker", "Value", "Unit", "Ref"], [[m.label, f"{m.value:g}", b.unit, m.citation or ""] for m in b.markers])
        if b.message:
            _para(doc, b.message, size=9.5)
    elif t == "conflict":
        _heading(doc, b.title or f"Documented values — {b.subject} / {b.parameter}", 2)
        rows = [[c.value, c.unit or "", c.document_id, c.revision or "", c.page or "", c.source, c.context or "", c.citation or "",
                 "preferred" if b.preferred_index == i else ""] for i, c in enumerate(b.claims)]
        _table(doc, ["Value", "Unit", "Document", "Rev", "Page", "Source", "Context", "Ref", ""], rows)
        if b.resolution:
            _para(doc, b.resolution, size=9.5)
    elif t == "graph":
        _heading(doc, b.title or "Topology", 2)
        names = {n.id: n.label for n in b.nodes}
        _table(doc, ["From", "Relation", "To", "Ref", "Inferred"],
               [[names.get(e.source, e.source), e.label, names.get(e.target, e.target), e.citation or "", "yes" if e.inferred else ""] for e in b.edges])
    elif t == "clarification":
        _heading(doc, "Clarification needed", 2)
        _para(doc, b.question)
        for m in b.missing:
            _para(doc, f"missing: {m}", style="List Bullet")
    elif t == "image":
        _para(doc, f"[figure: {b.caption or b.url}]", italic=True)


def _md_para(doc, style: str, text: str) -> None:
    if style.startswith("heading"):
        _heading(doc, text, min(3, int(style[-1]) + 1))
    elif style == "bullet":
        _para(doc, text, style="List Bullet")
    elif style == "number":
        _para(doc, text, style="List Number")
    elif style == "table_row":
        _para(doc, text, size=8)
    else:
        _para(doc, text)


def build_docx(resp, spec: DeliverableSpec, out_path: Path, figures=None) -> Path:
    doc = Document()
    doc.core_properties.author = spec.author
    doc.core_properties.title = spec.title
    doc.core_properties.category = spec.classification or ""
    doc.core_properties.subject = spec.status_line
    stamp = f"{spec.classification or 'UNCLASSIFIED'} · {spec.status_line}"
    for section in doc.sections:
        hp = section.header.paragraphs[0]
        hp.text = stamp
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in hp.runs:
            r.font.size = Pt(8)
            r.font.color.rgb = RED if not spec.signed_off else GREY
            r.font.bold = True
        fp = section.footer.paragraphs[0]
        fp.text = f"{stamp} · {spec.author} · {spec.generated_at}"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in fp.runs:
            r.font.size = Pt(8)
            r.font.color.rgb = GREY

    doc.add_heading(spec.title, 0)
    if spec.subtitle:
        _para(doc, spec.subtitle, italic=True, color=GREY)
    if not spec.signed_off:
        _para(doc, f"{spec.status_line}: every flagged figure must be resolved by a reviewer before this document is final. "
                   f"Nothing in it has been approved automatically.", bold=True, color=RED, size=10)
    else:
        _para(doc, spec.status_line, bold=True, color=GREY, size=10)
    if spec.classification:
        _para(doc, f"Classification: {spec.classification}", bold=True, size=9)

    _heading(doc, "Answer", 1)
    for style, text in markdown_to_paragraphs(strip_banner(resp.answer_markdown)):
        _md_para(doc, style, text)

    body_blocks = [b for b in resp.blocks if b.type not in REASONING_BLOCKS and b.type != "evidence"]
    if body_blocks:
        _heading(doc, "Supporting material", 1)
        for b in body_blocks:
            _render_block(doc, b)

    steps = calculation_steps_from_response(resp)
    if steps:
        _heading(doc, "Calculation steps", 1)
        _para(doc, "Every step is shown so the arithmetic can be repeated by hand; references point at the evidence list.", italic=True, size=9, color=GREY)
        _table(doc, ["Step", "Expression", "Result", "Note", "Ref"],
               [[s["step"], s["expression"], s["result"], s["note"], ", ".join(r for r in s["refs"] if r)] for s in steps])

    figs = figures if figures is not None else figures_from_response(resp)
    _heading(doc, "Provenance", 1)
    _para(doc, "Each figure links to the exact document, page and source line it was read from.", italic=True, size=9, color=GREY)
    rows = []
    for f in figs:
        d = f if isinstance(f, dict) else f.model_dump()
        src = d.get("source") or {}
        rows.append([d.get("label", ""), d.get("value", ""), d.get("unit") or "",
                     d.get("document_id") or src.get("document_id") or "—",
                     d.get("page") or src.get("page") or "—",
                     (d.get("source_text") or src.get("text") or "")[:120],
                     ("FLAGGED: " + (d.get("flag_reason") or "")) if d.get("flagged") else (d.get("citation") or "")])
    _table(doc, ["Figure", "Value", "Unit", "Document", "Page", "Source excerpt", "Ref / flag"], rows or [["—", "", "", "", "", "no figures", ""]])

    _heading(doc, "Evidence", 1)
    for ev in evidence_rows(resp):
        p = doc.add_paragraph()
        r = p.add_run(f"{ev['ref']} {ev['document_id']} p.{ev['page'] if ev['page'] is not None else '?'} ({ev['source']}): ")
        r.bold = True
        r.font.size = Pt(9)
        p.add_run(ev["text"][:600]).font.size = Pt(9)

    _heading(doc, "Audit", 1)
    _para(doc, audit_line(resp) + f" · model {getattr(resp, 'backend', '')}", size=8, color=GREY)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
