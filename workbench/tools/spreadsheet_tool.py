"""spreadsheet_read / spreadsheet_write over .xlsx files in the workspace (openpyxl).

Formulas are kept as formulas: a cell written as ``"=B2*1.1"`` stays a formula in the file, and
``spreadsheet_read`` reports both the formula text and, when the workbook has cached values, the
value. Nothing here evaluates formulas — that is the spreadsheet application's job, and the
calculate tool's when the arithmetic has to be shown step by step.
"""
from __future__ import annotations

import re

from workbench.tools.base import Tool, ToolContext, ToolError, ToolResult, confine

_CELL = re.compile(r"^[A-Za-z]{1,3}[1-9][0-9]{0,6}$")


def _load(path, *, formulas: bool):
    from openpyxl import load_workbook

    return load_workbook(path, data_only=not formulas)


class SpreadsheetReadTool(Tool):
    name = "spreadsheet_read"
    description = "Read a sheet of an .xlsx workbook in the workspace: rows of values, plus every formula found."
    parameters = {"type": "object",
                  "properties": {"path": {"type": "string"}, "sheet": {"type": "string"},
                                 "max_rows": {"type": "integer"}},
                  "required": ["path"]}

    def run(self, args: dict, ctx: ToolContext) -> ToolResult:
        target = confine(ctx.workspace, args["path"])
        if not target.exists():
            raise ToolError(f"no such workbook in the workspace: {args['path']!r}")
        max_rows = int(args.get("max_rows") or 200)
        wb_f = _load(target, formulas=True)
        wb_v = _load(target, formulas=False)
        name = args.get("sheet") or wb_f.sheetnames[0]
        if name not in wb_f.sheetnames:
            raise ToolError(f"no sheet {name!r}; sheets are {wb_f.sheetnames}")
        ws_f, ws_v = wb_f[name], wb_v[name]
        rows: list[list] = []
        formulas: list[dict] = []
        for r_idx, (row_f, row_v) in enumerate(zip(ws_f.iter_rows(), ws_v.iter_rows()), start=1):
            if r_idx > max_rows:
                break
            out_row = []
            for c_f, c_v in zip(row_f, row_v):
                val = c_f.value
                if isinstance(val, str) and val.startswith("="):
                    formulas.append({"cell": c_f.coordinate, "formula": val, "cached_value": c_v.value})
                    out_row.append(c_v.value if c_v.value is not None else val)
                else:
                    out_row.append(val)
            rows.append(out_row)
        lines = [f"sheet {name!r}: {ws_f.max_row} rows x {ws_f.max_column} columns (showing {len(rows)})"]
        for r in rows[:25]:
            lines.append(" | ".join("" if v is None else str(v) for v in r))
        if formulas:
            lines.append("formulas: " + "; ".join(f"{f['cell']} {f['formula']}" for f in formulas[:20]))
        return ToolResult(tool=self.name, output="\n".join(lines),
                          data={"path": args["path"], "sheet": name, "sheets": wb_f.sheetnames, "rows": rows,
                                "formulas": formulas, "dimensions": {"rows": ws_f.max_row, "columns": ws_f.max_column}})


class SpreadsheetWriteTool(Tool):
    name = "spreadsheet_write"
    description = ("Create or edit an .xlsx workbook in the workspace. Give either 'cells' ([{cell:'B3', value:...}]) "
                   "for targeted edits, 'rows' (list of lists, written from A1 or appended with append=true). "
                   "Strings starting with '=' are written as formulas.")
    parameters = {"type": "object",
                  "properties": {"path": {"type": "string"}, "sheet": {"type": "string"},
                                 "cells": {"type": "array"}, "rows": {"type": "array"}, "append": {"type": "boolean"},
                                 "header": {"type": "array"}},
                  "required": ["path"]}

    def run(self, args: dict, ctx: ToolContext) -> ToolResult:
        from openpyxl import Workbook, load_workbook

        target = confine(ctx.workspace, args["path"])
        if target.suffix.lower() != ".xlsx":
            raise ToolError("spreadsheet_write only writes .xlsx files")
        if target.exists():
            wb = load_workbook(target)
        else:
            wb = Workbook()
            wb.remove(wb.active)
        name = args.get("sheet") or (wb.sheetnames[0] if wb.sheetnames else "Sheet1")
        ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
        changed: list[dict] = []
        header = args.get("header")
        if header and ws.max_row == 1 and ws.cell(1, 1).value is None:
            for c, v in enumerate(header, start=1):
                ws.cell(1, c, v)
                changed.append({"cell": ws.cell(1, c).coordinate, "value": v})
        for item in args.get("cells") or []:
            if not isinstance(item, dict) or not _CELL.match(str(item.get("cell", ""))):
                raise ToolError(f"bad cell edit {item!r}; expected {{'cell': 'B3', 'value': ...}}")
            ws[item["cell"]] = item.get("value")
            changed.append({"cell": item["cell"].upper(), "value": item.get("value")})
        rows = args.get("rows") or []
        if rows:
            if not all(isinstance(r, list) for r in rows):
                raise ToolError("rows must be a list of lists")
            start = (ws.max_row + 1) if (args.get("append") and ws.cell(1, 1).value is not None) else (2 if header else 1)
            if header and not args.get("append"):
                start = 2
            for i, row in enumerate(rows):
                for j, value in enumerate(row):
                    cell = ws.cell(start + i, j + 1, value)
                    changed.append({"cell": cell.coordinate, "value": value})
        target.parent.mkdir(parents=True, exist_ok=True)
        wb.save(target)
        rel = target.relative_to(ctx.workspace.resolve()).as_posix()
        formulas = [c for c in changed if isinstance(c["value"], str) and str(c["value"]).startswith("=")]
        return ToolResult(tool=self.name, files=[rel],
                          output=f"{rel} sheet {name!r}: {len(changed)} cell(s) written" + (f", {len(formulas)} formula(s)" if formulas else ""),
                          data={"path": rel, "sheet": name, "changed": changed, "formulas": formulas})
